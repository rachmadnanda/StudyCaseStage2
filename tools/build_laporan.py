#!/usr/bin/env python3
"""
Membangun laporan .docx dari docs/laporan.md.

Pakai:
    python tools/build_laporan.py
"""

import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "docs", "laporan.md")
OUT = os.path.join(ROOT, "docs", "Laporan-CO-Zero-SIC8.docx")

BODY_FONT = "Times New Roman"
HEAD_FONT = "Calibri"
CODE_FONT = "Courier New"
ACCENT = RGBColor(0x1A, 0x3C, 0x6E)

BOX_MAP = str.maketrans({
    "─": "-", "│": "|", "┌": "+", "┐": "+", "└": "+", "┘": "+",
    "├": "+", "┤": "+", "┬": "+", "┴": "+", "┼": "+",
    "▶": ">", "◀": "<", "▲": "^", "▼": "v", "━": "-", "┃": "|",
})


# --------------------------------------------------------------------------- #
# Helper
# --------------------------------------------------------------------------- #
def set_cell_bg(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def shade_paragraph(paragraph, hex_color):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    p_pr.append(shd)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*\n]+\*)")


def add_inline(paragraph, text, size=None, bold_all=False, color=None):
    """Menambahkan teks dengan dukungan **bold**, *italic*, dan `code`."""
    for token in INLINE_RE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = CODE_FONT
            run._element.rPr.rFonts.set(qn("w:eastAsia"), CODE_FONT)
            if size:
                run.font.size = Pt(size - 1)
            else:
                run.font.size = Pt(9.5)
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(token)

        if bold_all:
            run.bold = True
        if size and not (token.startswith("`") and token.endswith("`")):
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color


# --------------------------------------------------------------------------- #
# Document setup
# --------------------------------------------------------------------------- #
def new_document():
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    pf = style.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15

    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.8)
        section.right_margin = Cm(2.2)

    # footer
    footer_p = doc.sections[0].footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer_p.add_run("CO-Zero  |  Tim Iki jenenge wes wkwk  |  SIC Batch 8  —  halaman ")
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    add_page_number(footer_p)
    for r in footer_p.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    return doc


def add_toc(doc, headings):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DAFTAR ISI")
    r.bold = True
    r.font.name = HEAD_FONT
    r.font.size = Pt(14)
    r.font.color.rgb = ACCENT
    p.paragraph_format.space_after = Pt(12)

    for level, text in headings:
        if level > 2:
            continue
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(3)
        para.paragraph_format.line_spacing = 1.1
        if level == 2:
            para.paragraph_format.left_indent = Cm(0.8)
        run = para.add_run(text)
        run.font.size = Pt(11)
        run.bold = (level == 1)

    doc.add_page_break()


def build_cover(doc, lines):
    for _ in range(3):
        doc.add_paragraph()

    title = lines[0].lstrip("# ").strip()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(20)
    r.font.name = HEAD_FONT
    r.font.color.rgb = ACCENT

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Laporan Proyek IoT — Stage 2")
    r.font.size = Pt(13)
    r.italic = True

    for _ in range(2):
        doc.add_paragraph()

    for line in lines[1:]:
        line = line.strip()
        if not line:
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, line, size=12)
        p.paragraph_format.space_after = Pt(2)

    for _ in range(4):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("POLITEKNIK NEGERI MALANG")
    r.bold = True
    r.font.size = Pt(13)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("2026")
    r.bold = True
    r.font.size = Pt(12)

    doc.add_page_break()


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.name = HEAD_FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)
    r.font.color.rgb = ACCENT if level <= 2 else RGBColor(0x33, 0x33, 0x33)
    r.font.size = Pt({1: 17, 2: 14, 3: 12}.get(level, 11.5))
    return p


def add_code_block(doc, code_lines):
    text = "\n".join(code_lines).translate(BOX_MAP)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.2)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.0
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT   # ASCII art wajib rata kiri
    shade_paragraph(p, "F2F4F7")
    run = p.add_run(text)
    run.font.name = CODE_FONT
    rf = run._element.rPr.rFonts
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(attr), CODE_FONT)
    run.font.size = Pt(8)
    return p


def _repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)

    # Lebar kolom: proporsional terhadap panjang isi, tetapi setiap kolom
    # diberi batas minimum berdasarkan kata terpanjang di dalamnya agar tidak
    # terpotong di tengah kata ("Kompo|nen").
    raw = []
    min_cm = []
    for j in range(cols):
        texts = [re.sub(r"[*`]", "", r[j]) if j < len(r) else "" for r in rows]
        longest = max(len(t) for t in texts) or 1
        longest_word = max((len(w) for t in texts for w in t.split()), default=4)
        raw.append(max(3.0, float(longest)))
        min_cm.append(max(1.3, 0.20 * longest_word + 0.45))

    usable_cm = 16.0
    if sum(min_cm) > usable_cm:          # terlalu banyak kolom: bagi rata
        min_cm = [usable_cm / cols] * cols

    total = sum(raw)
    widths_cm = [max(min_cm[j], usable_cm * raw[j] / total) for j in range(cols)]

    # normalisasi agar total tepat 16 cm, tanpa mengecilkan di bawah minimum
    for _ in range(cols):
        sisa = usable_cm - sum(widths_cm)
        if abs(sisa) < 0.01:
            break
        if sisa > 0:
            total_w = sum(widths_cm)
            widths_cm = [w + sisa * w / total_w for w in widths_cm]
        else:
            longgar = [j for j in range(cols) if widths_cm[j] > min_cm[j] + 0.05]
            if not longgar:
                break
            total_l = sum(widths_cm[j] - min_cm[j] for j in longgar)
            for j in longgar:
                widths_cm[j] += sisa * (widths_cm[j] - min_cm[j]) / total_l

    widths = [Cm(w) for w in widths_cm]

    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for i, row in enumerate(rows):
        _no_split(table.rows[i])
        if i == 0:
            _repeat_header(table.rows[i])
        for j in range(cols):
            cell = table.cell(i, j)
            cell.width = widths[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            value = row[j] if j < len(row) else ""
            add_inline(p, value, size=9, bold_all=(i == 0))
            if i == 0:
                set_cell_bg(cell, "1A3C6E")
                for run in p.runs:
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            elif i % 2 == 0:
                set_cell_bg(cell, "F4F6FA")

    # lebar grid eksplisit agar konsisten di Word & LibreOffice
    tbl_grid = table._tbl.find(qn("w:tblGrid"))
    if tbl_grid is not None:
        for gc, w in zip(tbl_grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(w.twips)))

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = 1.0


def add_bullet(doc, text, ordered=False):
    if ordered:
        return None
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    # indent eksplisit agar baris lanjutan sejajar dengan teks, bukan lebih menjorok
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    add_inline(p, text, size=11)


def add_numbered(doc, number, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.9)
    p.paragraph_format.first_line_indent = Cm(-0.6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(f"{number}.\t")
    r.font.size = Pt(11)
    add_inline(p, text, size=11)
    return p


def add_quote(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.1
    shade_paragraph(p, "FFF8E1")
    add_inline(p, text, size=10)
    for run in p.runs:
        run.italic = True


# --------------------------------------------------------------------------- #
# Markdown renderer
# --------------------------------------------------------------------------- #
def render_markdown(doc, md_lines):
    in_code = False
    code_buf = []
    table_buf = []
    para_buf = []
    quote_buf = []
    num = 0
    last_num_para = [None]

    def flush_paragraph():
        nonlocal para_buf
        if para_buf:
            text = " ".join(x.strip() for x in para_buf).strip()
            if text:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                add_inline(p, text)
            para_buf = []

    def flush_table():
        nonlocal table_buf
        if table_buf:
            rows = []
            for line in table_buf:
                cells = [c.strip() for c in line.strip().strip("|").split("|")]
                if all(re.fullmatch(r":?-{2,}:?", c or "-") for c in cells):
                    continue
                rows.append(cells)
            add_table(doc, rows)
            table_buf = []

    def flush_quote():
        nonlocal quote_buf
        if quote_buf:
            add_quote(doc, " ".join(x.strip() for x in quote_buf).strip())
            quote_buf = []

    for raw in md_lines:
        line = raw.rstrip("\n")

        if line.strip().startswith("```"):
            flush_paragraph()
            flush_table()
            if in_code:
                add_code_block(doc, code_buf)
                code_buf = []
                in_code = False
            else:
                in_code = True
            continue

        if in_code:
            code_buf.append(line)
            continue

        if line.strip().startswith("|"):
            flush_paragraph()
            table_buf.append(line)
            continue
        else:
            flush_table()

        if not line.strip():
            flush_paragraph()
            flush_quote()
            num = 0
            last_num_para[0] = None
            continue

        if line.strip() == "---":
            flush_paragraph()
            flush_quote()
            num = 0
            last_num_para[0] = None
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            flush_paragraph()
            flush_quote()
            num = 0
            last_num_para[0] = None
            add_heading(doc, m.group(2).strip(), len(m.group(1)))
            continue

        if line.lstrip().startswith("- "):
            flush_paragraph()
            flush_quote()
            num = 0
            last_num_para[0] = None
            add_bullet(doc, line.lstrip()[2:].strip())
            continue

        m = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if m:
            flush_paragraph()
            flush_quote()
            num += 1
            last_num_para[0] = add_numbered(doc, num, m.group(1).strip())
            continue

        if line.lstrip().startswith("> "):
            flush_paragraph()
            num = 0
            last_num_para[0] = None
            quote_buf.append(line.lstrip()[2:].strip())
            continue
        else:
            flush_quote()

        # baris lanjutan dari item bernomor
        if num > 0 and last_num_para[0] is not None and not para_buf:
            add_inline(last_num_para[0], " " + line.strip(), size=11)
            continue

        num = 0
        last_num_para[0] = None
        para_buf.append(line)

    flush_paragraph()
    flush_table()
    flush_quote()
    if code_buf:
        add_code_block(doc, code_buf)


def main():
    with open(SRC, encoding="utf-8") as f:
        lines = f.read().splitlines()

    # pisahkan cover (sebelum '---' pertama) dari isi
    split_at = next((i for i, l in enumerate(lines) if l.strip() == "---"), None)
    if split_at is None:
        cover_lines, body = [], lines
    else:
        cover_lines = [l for l in lines[:split_at] if l.strip()]
        body = lines[split_at + 1:]

    doc = new_document()
    if cover_lines:
        build_cover(doc, cover_lines)

    headings = []
    for line in body:
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            headings.append((len(m.group(1)), m.group(2).strip()))
    add_toc(doc, headings)

    render_markdown(doc, body)

    doc.save(OUT)
    print("Laporan dibuat:", OUT)


if __name__ == "__main__":
    main()
